# -*- coding: utf-8 -*-
"""
cli.py
======
Command-line interface for the Pashto stemmer.

Examples
--------
    # stem words given as arguments
    python -m pashto_stemmer.cli کورونه خبرونه افغانستان

    # stem a whole file (one token per output line, tab-separated)
    python -m pashto_stemmer.cli --file article.txt

    # read stdin, show the decision trace
    echo "د کورونو خبرونه" | python -m pashto_stemmer.cli --trace
"""

from __future__ import annotations

import argparse
import csv
import os
import sys

from .stemmer import PashtoStemmer, StemmerConfig


def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(
        prog="pashto-stem",
        description="A grammar-driven rule-based Pashto stemmer.",
    )
    p.add_argument("words", nargs="*", help="words to stem")
    p.add_argument("-f", "--file", help="stem all tokens in this UTF-8 text file")
    p.add_argument("-c", "--csv", "--excel", "--table", dest="table", metavar="PATH",
                   help="stem a column of this CSV, TSV or .xlsx workbook and "
                        "write a new column beside it")
    p.add_argument("--column", default="text", metavar="NAME",
                   help="which column to stem (default: text)")
    p.add_argument("--sheet", metavar="NAME",
                   help="which worksheet of an .xlsx file (default: the first)")
    p.add_argument("--out", metavar="PATH",
                   help="where to write the result (default: stdout)")
    p.add_argument("-t", "--trace", action="store_true",
                   help="print confidence and decision trace")
    p.add_argument("--no-dict", action="store_true", help="disable the lexicon")
    p.add_argument("--no-exceptions", action="store_true",
                   help="disable stopword/irregular/proper-noun handling")
    p.add_argument("--no-pos", action="store_true", help="disable POS filtering")
    p.add_argument("--no-compound", action="store_true",
                   help="disable compound decomposition")
    p.add_argument("--pos", action="store_true",
                   help="enable the POS filter (off by default, as in the library)")
    return p


def _read_table(path, sheet):
    """Return (fieldnames, rows) from a CSV, TSV or .xlsx file.

    CSV is read as utf-8-sig: a file exported from Excel begins with a
    byte-order mark, and without this the first column name comes back with the
    mark attached and the lookup fails.
    """
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise SystemExit("reading .xlsx needs openpyxl: pip install openpyxl")
        wb = load_workbook(path, data_only=True, read_only=True)
        ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
        it = ws.iter_rows(values_only=True)
        header = ["" if c is None else str(c) for c in next(it, ())]
        rows = [dict(zip(header, ["" if c is None else str(c) for c in r]))
                for r in it]
        wb.close()
        return header, rows
    delim = "\t" if ext in (".tsv", ".tab") else ","
    with open(path, encoding="utf-8-sig", newline="") as fh:
        reader = csv.DictReader(fh, delimiter=delim)
        return reader.fieldnames, list(reader)


def _write_table(path, fieldnames, rows):
    """Write to .xlsx when that is what was asked for, otherwise CSV/TSV."""
    ext = os.path.splitext(path)[1].lower() if path else ""
    if ext in (".xlsx", ".xlsm"):
        from openpyxl import Workbook
        wb = Workbook(); ws = wb.active
        ws.append(list(fieldnames))
        for r in rows:
            ws.append([r.get(k, "") for k in fieldnames])
        wb.save(path)
        return
    delim = "\t" if ext in (".tsv", ".tab") else ","
    out = open(path, "w", encoding="utf-8", newline="") if path else sys.stdout
    try:
        w = csv.DictWriter(out, fieldnames=list(fieldnames), delimiter=delim)
        w.writeheader()
        w.writerows(rows)
    finally:
        if path:
            out.close()


def stem_table(stemmer, path, column, out_path, sheet=None):
    """Stem one column of a table, writing the original rows back with a new
    column alongside. Works on CSV, TSV and Excel workbooks."""
    fieldnames, rows = _read_table(path, sheet)
    if not fieldnames:
        print(f"{path}: no header row", file=sys.stderr)
        return 1
    if column not in fieldnames:
        print(f"{path}: no column {column!r}; found {', '.join(fieldnames)}",
              file=sys.stderr)
        return 1

    field = column + "_stemmed"
    for row in rows:
        text = row.get(column) or ""
        row[field] = " ".join(r.stem for r in stemmer.stem_text(text))
    _write_table(out_path, list(fieldnames) + [field], rows)
    if out_path:
        print(f"{len(rows)} rows -> {out_path} (new column: {field})",
              file=sys.stderr)
    return 0


def main(argv=None) -> int:
    args = build_parser().parse_args(argv)
    # The defaults here must match StemmerConfig, or the command line and the
    # library give different answers for the same word. use_pos was defaulting
    # to True here and False there, so ښوونځی came out as ښوونځی from the CLI
    # and ښوون from the API.
    cfg = StemmerConfig(
        use_dictionary=not args.no_dict,
        use_exceptions=not args.no_exceptions,
        use_pos=args.pos,
        use_compound=not args.no_compound,
    )
    stemmer = PashtoStemmer(cfg)

    if args.table:
        return stem_table(stemmer, args.table, args.column, args.out, args.sheet)

    if args.file:
        with open(args.file, encoding="utf-8") as fh:
            text = fh.read()
        results = stemmer.stem_text(text)
    elif args.words:
        results = [stemmer.stem_word(w) for w in args.words]
    else:
        text = sys.stdin.read()
        results = stemmer.stem_text(text)

    out = open(args.out, "w", encoding="utf-8") if args.out else sys.stdout
    for r in results:
        if args.trace:
            print(f"{r.input}\t{r.stem}\t{r.confidence:.2f}\t{r.trace}", file=out)
        else:
            print(f"{r.input}\t{r.stem}", file=out)
    if args.out:
        out.close()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

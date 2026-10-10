# -*- coding: utf-8 -*-
"""
stem_file tests. Run: python tests/test_files.py

Covers the file API end to end: every parameter, both the streaming and the
in-memory path, and the failure modes a user will actually hit.
"""
import csv
import io
import os
import sys
import tempfile
from contextlib import redirect_stderr

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from pashto_stemmer import PashtoStemmer  # noqa: E402

ST = PashtoStemmer()


def _csv(body, header="word,freq\n", suffix=".csv"):
    fh = tempfile.NamedTemporaryFile("w", suffix=suffix, delete=False,
                                     encoding="utf-8", newline="")
    fh.write(header + body)
    fh.close()
    return fh.name


def test_returns_rows_and_writes_nothing_without_out():
    p = _csv("کورونه,10\nخبرونه,8\n")
    before = os.listdir(os.path.dirname(p))
    rows = ST.stem_file(p, column="word")
    after = os.listdir(os.path.dirname(p))
    os.unlink(p)
    assert rows[0]["word_stemmed"] == "کور"
    assert rows[1]["word_stemmed"] == "خبر"
    assert rows[0]["word"] == "کورونه" and rows[0]["freq"] == "10"
    assert len(before) == len(after), "nothing should have been written"


def test_writes_when_out_is_given():
    p = _csv("کورونه,10\n")
    o = p.replace(".csv", "_out.csv")
    info = ST.stem_file(p, column="word", out=o)
    rows = list(csv.DictReader(open(o, encoding="utf-8")))
    os.unlink(p); os.unlink(o)
    assert info["rows"] == 1 and info["column"] == "word_stemmed"
    assert rows[0]["word_stemmed"] == "کور"
    assert rows[0]["freq"] == "10", "other columns must survive"


def test_keep_original_false_drops_everything_else():
    p = _csv("کورونه,10\n")
    rows = ST.stem_file(p, column="word", keep_original=False)
    os.unlink(p)
    assert list(rows[0].keys()) == ["word_stemmed"]


def test_new_column_names_the_output():
    p = _csv("کورونه,10\n")
    rows = ST.stem_file(p, column="word", new_column="stem")
    os.unlink(p)
    assert rows[0]["stem"] == "کور" and "word_stemmed" not in rows[0]


def test_trace_adds_the_rules_that_fired():
    p = _csv("کورونه,10\n")
    rows = ST.stem_file(p, column="word", trace=True)
    os.unlink(p)
    assert rows[0]["word_rules"] == "+ونه", rows[0]


def test_unique_gives_the_same_answer_as_the_plain_path():
    p = _csv("کورونه,1\nخبرونه,2\nکورونه,3\n")
    plain = ST.stem_file(p, column="word")
    uniq = ST.stem_file(p, column="word", unique=True)
    os.unlink(p)
    assert [r["word_stemmed"] for r in plain] == [r["word_stemmed"] for r in uniq]


def test_missing_column_raises_and_names_the_real_columns():
    p = _csv("کورونه,10\n")
    try:
        ST.stem_file(p, column="nope")
        raised = None
    except KeyError as e:
        raised = str(e)
    os.unlink(p)
    assert raised and "nope" in raised and "word" in raised and "freq" in raised


def test_multiword_cells_warn_once():
    p = _csv("د کورونو خبرونه\n", header="text\n")
    err = io.StringIO()
    with redirect_stderr(err):
        rows = ST.stem_file(p, column="text")
    os.unlink(p)
    # the warning must name the offending cell, not just say "multi-word":
    # the earlier wording told the user to "pass the word column", which they
    # already had.
    msg = err.getvalue()
    assert "more than one word" in msg, msg
    assert "د کورونو خبرونه" in msg, msg
    assert rows[0]["text_stemmed"] == "د کور خبر"


def test_blank_cells_do_not_stop_the_run():
    p = _csv("کورونه,1\n   ,2\n", header="word,freq\n")
    rows = ST.stem_file(p, column="word")
    os.unlink(p)
    assert rows[0]["word_stemmed"] == "کور"
    assert rows[-1]["word_stemmed"] == ""


def test_tsv_is_read_with_tabs():
    p = _csv("کورونه\t1\n", header="word\tfreq\n", suffix=".tsv")
    rows = ST.stem_file(p, column="word")
    os.unlink(p)
    assert rows[0]["word_stemmed"] == "کور"


def test_streaming_and_in_memory_agree():
    p = _csv("".join(f"کورونه,{i}\nخبرونه,{i}\n" for i in range(50)))
    o = p.replace(".csv", "_out.csv")
    memory = ST.stem_file(p, column="word")
    ST.stem_file(p, column="word", out=o)
    streamed = list(csv.DictReader(open(o, encoding="utf-8")))
    os.unlink(p); os.unlink(o)
    assert [r["word_stemmed"] for r in memory] == \
           [r["word_stemmed"] for r in streamed]


def test_excel_round_trip():
    try:
        from openpyxl import Workbook, load_workbook
    except ImportError:
        print("  SKIP  openpyxl not installed")
        return
    wb = Workbook(); ws = wb.active
    ws.append(["word", "freq"]); ws.append(["کورونه", 10])
    src = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False); src.close()
    wb.save(src.name)
    dst = src.name.replace(".xlsx", "_out.xlsx")
    ST.stem_file(src.name, column="word", out=dst)
    rows = list(load_workbook(dst).active.iter_rows(values_only=True))
    os.unlink(src.name); os.unlink(dst)
    assert rows[0] == ("word", "freq", "word_stemmed")
    assert rows[1][2] == "کور"


def test_excel_sheet_can_be_chosen():
    try:
        from openpyxl import Workbook
    except ImportError:
        print("  SKIP  openpyxl not installed")
        return
    wb = Workbook(); wb.active.append(["word"]); wb.active.append(["کورونه"])
    other = wb.create_sheet("Other")
    other.append(["word"]); other.append(["پوهنتون"])
    src = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False); src.close()
    wb.save(src.name)
    rows = ST.stem_file(src.name, column="word", sheet="Other")
    os.unlink(src.name)
    assert rows[0]["word_stemmed"] == "پوهن"


def test_stemming_is_deterministic():
    """The same input gives the same answer every time.

    Note this is not idempotence: stem(stem(w)) can differ from stem(w),
    because the stemmer makes a single pass by design (max_passes = 1) and a
    stem may itself end in something the inventory recognises. On the
    development words where the two differ, the reference agrees with one pass
    62 times and with two passes 58 times, so neither reading is clearly right
    and the single pass is a documented choice, not an oversight.
    """
    st = PashtoStemmer()
    for w in ("کورونه", "خدمتګارانو", "ښوونځیو", "اړتیاو"):
        first = st.stem(w)
        st.clear_cache()
        assert st.stem(w) == first, w
        assert st.stem(w) == first, w


def test_a_renamed_csv_reports_itself_clearly():
    """openpyxl says "BadZipFile: File is not a zip file", which helps nobody.

    A CSV renamed to .xlsx is a common mistake; the message has to name it.
    """
    import pashto_stemmer.files as F
    p = os.path.join(tempfile.mkdtemp(), "words.xlsx")
    with open(p, "w", encoding="utf-8") as fh:
        fh.write("word\nکورونه\n")
    try:
        F.read_rows(p)
    except SystemExit as exc:
        msg = str(exc)
        assert "not a readable Excel workbook" in msg, msg
        assert "plain text" in msg, msg
    else:
        raise AssertionError("a text file named .xlsx should have been refused")


def _run_all():
    fns = [g for n, g in sorted(globals().items()) if n.startswith("test_")]
    passed = 0
    for fn in fns:
        try:
            fn(); print(f"  PASS  {fn.__name__}"); passed += 1
        except AssertionError as e:
            print(f"  FAIL  {fn.__name__}: {e}")
        except Exception as e:
            print(f"  ERROR {fn.__name__}: {type(e).__name__}: {e}")
    print(f"\n{passed}/{len(fns)} passed")
    return passed == len(fns)


if __name__ == "__main__":
    sys.exit(0 if _run_all() else 1)

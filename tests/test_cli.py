# -*- coding: utf-8 -*-
"""
Command-line interface tests. Run: python tests/test_cli.py

The CLI is what most people will actually touch, so the text and CSV modes are
covered here rather than left to manual checking.
"""
import csv
import io
import os
import sys
import tempfile
from contextlib import redirect_stdout, redirect_stderr

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from pashto_stemmer.cli import main  # noqa: E402


def _run(argv):
    out, err = io.StringIO(), io.StringIO()
    with redirect_stdout(out), redirect_stderr(err):
        code = main(argv)
    return code, out.getvalue(), err.getvalue()


def _csv(rows, header="id,text\n", bom=False):
    fh = tempfile.NamedTemporaryFile("w", suffix=".csv", delete=False,
                                     encoding="utf-8-sig" if bom else "utf-8",
                                     newline="")
    fh.write(header + rows)
    fh.close()
    return fh.name


def test_words_on_the_command_line():
    code, out, _ = _run(["کورونه", "افغانستان"])
    assert code == 0
    lines = out.strip().split("\n")
    assert lines[0].split("\t") == ["کورونه", "کور"], lines[0]
    assert lines[1].split("\t") == ["افغانستان", "افغان"], lines[1]


def test_text_file():
    fh = tempfile.NamedTemporaryFile("w", suffix=".txt", delete=False,
                                     encoding="utf-8")
    fh.write("د کورونو خبرونه\n"); fh.close()
    code, out, _ = _run(["--file", fh.name])
    os.unlink(fh.name)
    assert code == 0
    assert [l.split("\t")[1] for l in out.strip().split("\n")] == ["د", "کور", "خبر"]


def test_csv_adds_a_column_and_keeps_the_originals():
    path = _csv("1,د کورونو خبرونه\n2,پوهنتون ته ولاړ\n")
    code, out, _ = _run(["--csv", path, "--column", "text"])
    os.unlink(path)
    assert code == 0
    rows = list(csv.DictReader(io.StringIO(out)))
    assert list(rows[0].keys()) == ["id", "text", "text_stemmed"]
    assert rows[0]["text"] == "د کورونو خبرونه"      # original is untouched
    assert rows[0]["text_stemmed"] == "د کور خبر"
    assert rows[1]["text_stemmed"] == "پوهن ته ولاړ"


def test_csv_writes_to_a_file():
    path = _csv("1,کورونه\n")
    out_path = path.replace(".csv", "_out.csv")
    code, _, err = _run(["--csv", path, "--column", "text", "--out", out_path])
    assert code == 0
    rows = list(csv.DictReader(open(out_path, encoding="utf-8")))
    os.unlink(path); os.unlink(out_path)
    assert rows[0]["text_stemmed"] == "کور"
    assert "1 rows" in err                            # progress goes to stderr


def test_csv_reports_a_missing_column_instead_of_writing_rubbish():
    path = _csv("1,کورونه\n")
    code, out, err = _run(["--csv", path, "--column", "body"])
    os.unlink(path)
    assert code == 1
    assert "no column 'body'" in err and "id, text" in err
    assert out == ""                                  # nothing was written


def test_csv_handles_an_excel_byte_order_mark():
    path = _csv("1,کورونه\n", bom=True)
    code, out, _ = _run(["--csv", path, "--column", "text"])
    os.unlink(path)
    assert code == 0
    rows = list(csv.DictReader(io.StringIO(out)))
    assert rows[0]["text_stemmed"] == "کور"


def test_excel_workbook_round_trip():
    try:
        from openpyxl import Workbook, load_workbook
    except ImportError:
        print("  SKIP  openpyxl not installed")
        return
    wb = Workbook(); ws = wb.active
    ws.append(["id", "text"]); ws.append([1, "د کورونو خبرونه"])
    src = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False); src.close()
    wb.save(src.name)
    dst = src.name.replace(".xlsx", "_out.xlsx")

    code, _, _ = _run(["--excel", src.name, "--column", "text", "--out", dst])
    assert code == 0
    rows = list(load_workbook(dst).active.iter_rows(values_only=True))
    os.unlink(src.name); os.unlink(dst)
    assert rows[0] == ("id", "text", "text_stemmed"), rows[0]
    assert rows[1][2] == "د کور خبر", rows[1]


def test_excel_to_csv_on_stdout():
    try:
        from openpyxl import Workbook
    except ImportError:
        print("  SKIP  openpyxl not installed")
        return
    wb = Workbook(); ws = wb.active
    ws.append(["id", "text"]); ws.append([1, "کورونه"])
    src = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False); src.close()
    wb.save(src.name)
    code, out, _ = _run(["--excel", src.name, "--column", "text"])
    os.unlink(src.name)
    assert code == 0
    rows = list(csv.DictReader(io.StringIO(out)))
    assert rows[0]["text_stemmed"] == "کور"


def test_tsv_is_read_with_tabs():
    fh = tempfile.NamedTemporaryFile("w", suffix=".tsv", delete=False,
                                     encoding="utf-8", newline="")
    fh.write("id\ttext\n1\tکورونه\n"); fh.close()
    code, out, _ = _run(["--table", fh.name, "--column", "text"])
    os.unlink(fh.name)
    assert code == 0
    rows = list(csv.DictReader(io.StringIO(out)))
    assert rows[0]["text_stemmed"] == "کور"


def test_cli_and_library_agree():
    from pashto_stemmer import PashtoStemmer
    st = PashtoStemmer()
    for w in ("ښوونځی", "کورونو", "افغانستان", "پوهنتون"):
        _, out, _ = _run([w])
        assert out.strip().split("\t")[1] == st.stem(w), w


def test_cli_defaults_match_stemmer_config():
    """Every behaviour flag must leave StemmerConfig's own default in place.

    Two flags had drifted: --no-pos was parsed and never used while use_pos
    defaulted to True on the command line and False in the library, and
    --no-compound turned compound decomposition ON by its absence. The same
    word then came back differently from the CLI and the API.
    """
    import argparse
    from pashto_stemmer.cli import build_parser
    from pashto_stemmer.stemmer import StemmerConfig
    args = build_parser().parse_args([])
    d = StemmerConfig()
    assert args.column == "word", args.column
    assert args.max_passes == d.max_passes
    assert args.pos is d.use_pos
    assert args.compound is d.use_compound
    assert args.lemmatize is d.use_verb_dictionary
    for flag, field in (("no_dict", "use_dictionary"),
                        ("no_exceptions", "use_exceptions"),
                        ("no_prefixes", "use_prefixes"),
                        ("no_suffixes", "use_suffixes"),
                        ("no_derivational", "strip_derivational")):
        assert getattr(args, flag) is not getattr(d, field), flag


def test_every_behaviour_flag_is_wired():
    """A flag that parses but changes nothing is how the last bug got in."""
    cases = [(["--lemmatize", "شو"], "کېدل"),
             (["شو"], "شو"),
             (["--no-suffixes", "کورونه"], "کورونه"),
             (["--no-exceptions", "مخې"], "مخې"),
             (["مخې"], "مخ"),
             (["--no-prefixes", "همکار"], "همکار"),
             (["همکار"], "کار"),
             (["--no-derivational", "نیمګړی"], "نیمګړ"),
             (["--pos", "ښوونځی"], "ښوونځی")]
    for argv, expected in cases:
        _, out, _ = _run(argv)
        got = out.strip().split("\t")[1]
        assert got == expected, f"{argv} gave {got}, expected {expected}"


def test_deprecated_flags_are_still_accepted():
    """Scripts written against the earlier release must keep working."""
    code, out, _ = _run(["--no-compound", "--no-pos", "کورونه"])
    assert code == 0
    assert out.strip().split("\t")[1] == "کور"


def _run_all():
    fns = [g for n, g in globals().items() if n.startswith("test_")]
    passed = 0
    for fn in fns:
        try:
            fn(); print(f"  PASS  {fn.__name__}"); passed += 1
        except AssertionError as e:
            print(f"  FAIL  {fn.__name__}: {e}")
        except Exception as e:
            print(f"  ERROR {fn.__name__}: {type(e).__name__}: {e}")
    print(f"\n{passed}/{len(fns)} passed")
    return passed == len(fns)


if __name__ == "__main__":
    sys.exit(0 if _run_all() else 1)
